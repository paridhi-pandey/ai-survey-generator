"""
Excel Exporter — Generates a styled Excel workbook from the survey question list.
"""

from io import BytesIO
from typing import List, Dict

try:
    import openpyxl
    from openpyxl.styles import (
        Font, PatternFill, Alignment, Border, Side
    )
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


# Column definitions: (header label, dict key, width)
COLUMNS = [
    ("No.", "no", 6),
    ("Section", "section", 22),
    ("Session No.", "session_no", 12),
    ("Session", "session", 30),
    ("Category", "category", 18),
    ("Question", "question", 55),
    ("Question Type", "question_type", 16),
    ("Options", "options", 40),
    ("Required", "required", 10),
]

# Header styling
HEADER_FILL = "1E3A5F"      # Dark navy
HEADER_FONT_COLOR = "FFFFFF"
SECTION_FILLS = {
    "Participant Information": "D6E4F7",
    "Session Feedback": "E8F5E9",
    "Overall Workshop": "FFF9C4",
    "Suggestions": "FCE4EC",
}
ALT_FILL = "F7F9FC"


def export_to_excel(survey_rows: List[Dict]) -> bytes:
    """
    Generate a styled Excel workbook from the survey question list.

    Args:
        survey_rows: List of survey question dicts from build_full_survey().

    Returns:
        Excel file as bytes (ready for st.download_button).
    """
    if not OPENPYXL_AVAILABLE:
        raise ImportError("openpyxl not installed. Run: pip install openpyxl")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Feedback Form"

    # --- Write header row ---
    header_font = Font(bold=True, color=HEADER_FONT_COLOR, name="Calibri", size=11)
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    for col_idx, (header, _, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"

    # --- Write data rows ---
    for row_idx, row_data in enumerate(survey_rows, start=2):
        section = row_data.get("section", "")
        section_fill_color = SECTION_FILLS.get(section, ALT_FILL if row_idx % 2 == 0 else "FFFFFF")
        row_fill = PatternFill("solid", fgColor=section_fill_color)

        for col_idx, (_, key, _) in enumerate(COLUMNS, start=1):
            value = row_data.get(key, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=str(value) if value != "" else "")
            cell.fill = row_fill
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)
            cell.alignment = Alignment(
                horizontal="left" if col_idx not in (1, 3, 9) else "center",
                vertical="top",
                wrap_text=True,
            )

        ws.row_dimensions[row_idx].height = 18

    # --- Title row above header ---
    ws.insert_rows(1)
    ws.merge_cells(f"A1:{get_column_letter(len(COLUMNS))}1")
    title_cell = ws["A1"]
    title_cell.value = "Workshop Feedback Survey"
    title_cell.font = Font(bold=True, size=14, color="1E3A5F", name="Calibri")
    title_cell.fill = PatternFill("solid", fgColor="EBF2FA")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def read_xlsx_as_text(file_bytes: bytes) -> str:
    """
    Read an XLSX reference file and return its content as plain text.
    Used to pass reference survey style to Gemini.

    Args:
        file_bytes: Raw bytes of XLSX file.

    Returns:
        String representation of sheet content.
    """
    if not OPENPYXL_AVAILABLE:
        return ""

    try:
        wb = openpyxl.load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"Sheet: {ws.title}")
            for row in ws.iter_rows(max_row=50, values_only=True):
                row_text = " | ".join(str(c) if c is not None else "" for c in row)
                if row_text.strip():
                    lines.append(row_text)
        return "\n".join(lines)
    except Exception:
        return ""
